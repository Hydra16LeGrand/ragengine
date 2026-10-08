"""Garde-fou bombe de décompression `.xlsx`/`.xlsm` (audit sécurité 2026-09, constat 10).

Les archives pathologiques sont fabriquées à la main (`zipfile`) plutôt que par
`openpyxl` : c'est justement le format brut du conteneur, pas ce qu'un classeur
légitime peut produire, qui est visé ici. Le classeur « normal » de comparaison,
lui, vient de `tests.ingestion.precheck.documents.workbook_xlsx` — un vrai
fichier `openpyxl`, comme le reste de la suite precheck.
"""

from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest

from app.core.config import Settings
from app.ingestion.precheck.extraction import ExtractionError
from app.ingestion.precheck.spreadsheet import analyze_xlsx
from tests.ingestion.precheck import documents

_TWENTY_MIO = 20 * 1024 * 1024


def _bomb_zip(*, uncompressed_size: int, name: str = "xl/worksheets/sheet1.xml") -> bytes:
    """Archive ZIP à un seul membre, hautement compressible : simule une bombe réelle."""
    payload = b"0" * uncompressed_size
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(name, payload)
    return buffer.getvalue()


def _many_small_entries_zip(*, count: int, size_each: int) -> bytes:
    """Archive à de nombreuses petites entrées : chacune sous le seuil de ratio, mais la
    somme dépasse un plafond de taille totale bas."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for index in range(count):
            archive.writestr(f"xl/worksheets/sheet{index}.xml", "a" * size_each)
    return buffer.getvalue()


def test_a_normal_workbook_is_accepted() -> None:
    """Un classeur réel, même dense, passe le garde-fou sans y être ralenti."""
    content = documents.workbook_xlsx(
        {"Synthese": [("Poste", "Montant"), ("Achats", 12), ("Ventes", 34)]}
    )

    outcome = analyze_xlsx(content, Settings())

    assert outcome.table_count == 1


def test_a_workbook_with_an_absurd_ratio_is_refused_under_default_settings() -> None:
    """20 Mio de zéros compressent à quelques Kio : très au-delà du plafond par défaut
    (100:1) — refusé sans qu'il soit nécessaire d'abaisser le seuil pour le test."""
    bomb = _bomb_zip(uncompressed_size=_TWENTY_MIO)

    with pytest.raises(ExtractionError, match="ratio"):
        analyze_xlsx(bomb, Settings())


def test_openpyxl_never_opens_a_workbook_refused_by_the_guard(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """La preuve que le refus a lieu **avant** toute décompression réelle : si
    `openpyxl.load_workbook` était encore appelé, ce test échouerait."""
    import app.ingestion.precheck.spreadsheet as spreadsheet_module

    def _must_not_be_called(*args: object, **kwargs: object) -> None:
        raise AssertionError("openpyxl.load_workbook a été appelé malgré le refus du garde-fou")

    monkeypatch.setattr(spreadsheet_module.openpyxl, "load_workbook", _must_not_be_called)
    bomb = _bomb_zip(uncompressed_size=_TWENTY_MIO)

    with pytest.raises(ExtractionError):
        analyze_xlsx(bomb, Settings())


def test_a_workbook_beyond_the_total_decompressed_size_is_refused() -> None:
    """Plafond de taille totale, indépendant du ratio : de nombreuses petites entrées,
    chacune sous le seuil de vérification du ratio, dont la somme dépasse un plafond bas."""
    content = _many_small_entries_zip(count=50, size_each=100)
    settings = Settings(RAGENGINE_SPREADSHEET_MAX_DECOMPRESSED_BYTES=4000)

    with pytest.raises(ExtractionError, match="décompressée totale"):
        analyze_xlsx(content, settings)


def test_a_malformed_but_small_archive_fails_at_openpyxl_not_at_the_guard() -> None:
    """Une archive ZIP minuscule et non pathologique n'est jamais bloquée par le
    garde-fou : l'échec, s'il y en a un, vient d'`openpyxl` (contenu non-`.xlsx`),
    pas de la taille ou du ratio."""
    content = _bomb_zip(uncompressed_size=10)

    with pytest.raises(ExtractionError) as failure:
        analyze_xlsx(content, Settings())

    assert "ratio" not in str(failure.value)
    assert "décompressée totale" not in str(failure.value)


def test_the_precheck_reports_a_clean_message_not_a_raw_traceback() -> None:
    """Bout en bout via l'analyseur : le job échoue proprement (rapport dégradé), avec un
    message compréhensible, jamais une trace Python brute."""
    from app.ingestion.precheck.analyzer import analyze_document_sync
    from app.ingestion.precheck.schemas import PrecheckErrorCode

    bomb = _bomb_zip(uncompressed_size=_TWENTY_MIO)
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    report = analyze_document_sync(bomb, "classeur.xlsx", media_type)

    # `supported` porte sur le *format* (xlsx est bien un format connu), pas sur le
    # résultat de l'extraction : c'est `error` qui porte l'échec, dégradé mais propre.
    assert report.error is not None
    assert report.error.code is PrecheckErrorCode.EXTRACTION_FAILED
    assert "Traceback" not in report.error.message
    assert "ratio" in report.error.message


def test_a_real_repetitive_workbook_stays_under_the_default_thresholds() -> None:
    """Le cas le plus compressible qu'openpyxl produise réellement (mesuré : ~30:1, très
    sous le plafond par défaut de 100:1) n'est jamais confondu avec une bombe."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    for _ in range(2000):
        worksheet.append(["répétée"] * 10)
    buffer = io.BytesIO()
    workbook.save(buffer)

    outcome = analyze_xlsx(buffer.getvalue(), Settings())

    assert outcome.table_count == 1
